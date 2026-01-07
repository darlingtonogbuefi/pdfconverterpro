import os
from openpyxl import Workbook
from docx import Document
from pathlib import Path
from backend.utils.file_utils import encode_file_to_base64

def convert_word_to_excel(word_path: str, out_path: str, original_name: str):
    doc = Document(word_path)
    wb = Workbook()
    ws = wb.active
    row = 1

    for para in doc.paragraphs:
        if para.text.strip():
            ws.cell(row=row, column=1, value=para.text.strip())
            row += 1

    wb.save(out_path)

    return {
        "success": True,
        "filename": Path(original_name).stem + ".xlsx",
        "file": encode_file_to_base64(out_path)
    }
